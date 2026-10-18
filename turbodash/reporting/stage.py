"""Tables, console output and Word reports for the original stage model."""
import numpy as np


def _fmt(value, unit="-", width=10, prec=4):
    if value is None:
        return " " * width

    # Strings: right-aligned, same width as numbers, no unit
    if isinstance(value, str):
        return f"{value:>{width}s}"

    # Integers
    if isinstance(value, int):
        return f"{value:{width}d} {unit}".rstrip()

    # Floats
    if isinstance(value, float):
        return f"{value:{width}.{prec}f} {unit}".rstrip()

    # Fallback
    return f"{str(value):>{width}s}"


def _fmt_mm(value_m, width=10, prec=2):
    return _fmt(1e3 * value_m, "mm", width, prec)


def _fmt_deg(value, width=10, prec=2):
    return _fmt(value, "deg", width, prec)


def _section(title):
    print(f"\n{title}")
    print("=" * len(title))


def _pretty_name(key: str) -> str:
    return key.replace("_", " ").capitalize()


def _fmt_geom(key, value):
    """
    Geometry-aware formatter based on variable semantics.
    """
    if value is None:
        return " " * 12

    if key in {
        "radius_in",
        "radius_out",
        "height",
        "chord",
        "chord_true",
        "chord_meridional",
        "spacing",
        "opening",
        "maximum_thickness",
        "leading_edge_radius",
        "trailing_edge_thickness",
    }:
        return _fmt_mm(value)

    if key in {
        "flaring_angle",
        "trailing_edge_wedge_angle",
        "metal_angle_in",
        "metal_angle_out",
    }:
        return _fmt_deg(value)

    if key in {
        "aspect_ratio",
        "solidity",
        "flaring_angle",
        "trailing_edge_wedge_angle",
        "maximum_thickness_location",
    }:
        return _fmt(value, "-")

    if key == "blade_count":
        return _fmt(value, "-", prec=0)

    # Fallback
    return _fmt(value, "-")


def _print_geometry_block(title, geom):
    print(f"\n{title}:")
    for key, value in geom.items():
        label = _pretty_name(key)
        print(f"  {label:28s}: {_fmt_geom(key, value)}")


def print_stage(out):
    """
    ASCII-only, human-readable pretty print of turbine stage results.
    Assumes a fully nested and structured `out` dictionary.
    """

    inputs = out["inputs"]
    perf = out["stage_performance"]
    geom = out["geometry"]
    stations = out["flow_stations"]

    # ==============================================================
    # Boundary conditions and operating point
    # ==============================================================
    _section("Boundary conditions and operating point")

    print(f"{'Fluid name':30s}: {_fmt(inputs['fluid_name'])}")
    print(f"{'Stage type':30s}: {_fmt(inputs['stage_type'])}")
    print(f"{'Inlet pressure':30s}: {_fmt(stations[0]['p'] / 1e5, 'bar')}")
    print(f"{'Exit pressure':30s}: {_fmt(stations[-1]['p'] / 1e5, 'bar')}")
    print(f"{'Mass flow rate':30s}: {_fmt(inputs['mass_flow_rate'], 'kg/s')}")
    print(f"{'Rotational speed':30s}: {_fmt(perf['rotational_speed'], 'rpm', prec=1)}")

    # ==============================================================
    # Stage performance
    # ==============================================================
    _section("Stage performance")

    print(f"{'Total-to-total efficiency':30s}: {_fmt(perf['efficiency_tt'], '-')}")
    print(f"{'Total-to-static efficiency':30s}: {_fmt(perf['efficiency_ts'], '-')}")
    print(
        f"{'Total-to-static pressure ratio':30s}: {_fmt(perf['pressure_ratio_ts'], '-')}"
    )
    print(f"{'Total-to-static volume ratio':30s}: {_fmt(perf['volume_ratio_ts'], '-')}")
    print(f"{'Flow coefficient':30s}: {_fmt(perf['flow_coefficient'], '-')}")
    print(f"{'Loading coefficient':30s}: {_fmt(perf['work_coefficient'], '-')}")
    print(f"{'Degree of reaction':30s}: {_fmt(perf['degree_reaction'], '-')}")
    print(f"{'Blade velocity ratio':30s}: {_fmt(perf['blade_velocity_ratio'], '-')}")
    print(f"{'Specific speed':30s}: {_fmt(perf['specific_speed'], '-')}")
    print(f"{'Spouting velocity':30s}: {_fmt(perf['spouting_velocity'], 'm/s')}")
    print(f"{'Rotor exit blade speed':30s}: {_fmt(perf['rotor_exit_velocity'], 'm/s')}")

    # ==============================================================
    # Geometry summary
    # ==============================================================
    _section("Geometry summary")
    _print_geometry_block("Stator", geom["stator"])
    _print_geometry_block("Rotor", geom["rotor"])

    # ==============================================================
    # Flow stations
    # ==============================================================
    _section("Flow stations")
    header = (
        f"{'Stn':>3s} "
        f"{'p [bar]':>10s} "
        f"{'T [K]':>9s} "
        f"{'rho [kg/m3]':>12s} "
        f"{'v [m/s]':>10s} "
        f"{'w [m/s]':>10s} "
        f"{'Ma':>7s} "
        f"{'alpha [deg]':>12s} "
        f"{'beta [deg]':>11s} "
        f"{'r [mm]':>9s} "
        f"{'H [mm]':>9s}"
    )
    print()
    print("-" * len(header))

    print(header)

    for i, st in enumerate(stations):
        print(
            f"{i:3d} "
            f"{st['p']/1e5:10.3f} "
            f"{st['T']:9.2f} "
            f"{st['d']:12.4f} "
            f"{st['v']:10.2f} "
            f"{st['w']:10.2f} "
            f"{st['Ma']:7.3f} "
            f"{st['alpha']:12.2f} "
            f"{st['beta']:11.2f} "
            f"{1e3*st['r']:9.2f} "
            f"{1e3*st['H']:9.2f}"
        )
    print("-" * len(header))
    print()


def stage_performance_table(out):
    perf = out["stage_performance"]

    rows = [
        ("Total-to-total efficiency", perf["efficiency_tt"], "-"),
        ("Total-to-static efficiency", perf["efficiency_ts"], "-"),
        ("Pressure ratio (t-s)", perf["pressure_ratio_ts"], "-"),
        ("Volume ratio (t-s)", perf["volume_ratio_ts"], "-"),
        ("Flow coefficient", perf["flow_coefficient"], "-"),
        ("Work coefficient", perf["work_coefficient"], "-"),
        ("Degree of reaction", perf["degree_reaction"], "-"),
        ("Blade velocity ratio", perf["blade_velocity_ratio"], "-"),
        ("Specific speed", perf["specific_speed"], "-"),
        ("Spouting velocity", perf["spouting_velocity"], "m/s"),
        ("Rotor exit blade speed", perf["rotor_exit_velocity"], "m/s"),
        ("Rotational speed", perf["rotational_speed"], "rpm"),
        ("Isentropic power", perf["power_isentropic"] / 1e3, "kW"),
        ("Actual power (t-t)", perf["power_actual_tt"] / 1e3, "kW"),
        ("Actual power (t-s)", perf["power_actual_ts"] / 1e3, "kW"),
    ]

    return [
        {"Quantity": name, "Value": float(value), "Unit": unit}
        for name, value, unit in rows
    ]


def geometry_table(out):
    stator = out["geometry"]["stator"]
    rotor = out["geometry"]["rotor"]

    # Sanity check: geometry dictionaries must match
    if stator.keys() != rotor.keys():
        raise ValueError("Stator and rotor geometry keys do not match")

    def pretty_name(key):
        return key.replace("_", " ").capitalize()

    def infer_unit_and_scale(key):
        if key in {
            "radius_in",
            "radius_out",
            "height",
            "chord",
            "chord_true",
            "chord_meridional",
            "spacing",
            "opening",
            "maximum_thickness",
            "leading_edge_radius",
            "trailing_edge_thickness",
        }:
            return "mm", 1e3

        if key in {
            "metal_angle_in",
            "metal_angle_out",
            "flaring_angle",
            "trailing_edge_wedge_angle",
        }:
            return "deg", 1.0

        if key in {
            "solidity",
            "aspect_ratio",
            "maximum_thickness_location",
        }:
            return "-", 1.0

        if key == "blade_count":
            return "-", 1.0

        # Fallback
        return "-", 1.0

    rows = []
    for key in stator.keys():
        unit, scale = infer_unit_and_scale(key)

        rows.append(
            {
                "Variable": pretty_name(key),
                "Stator": float(stator[key] * scale),
                "Rotor": float(rotor[key] * scale),
                "Unit": unit,
            }
        )

    return rows


def flow_stations_table(out):
    rows = []

    for i, st in enumerate(out["flow_stations"]):
        rows.append(
            {
                "Station": i,
                "p [bar]": st["p"] / 1e5,
                "T [K]": st["T"],
                "ρ [kg/m³]": st["d"],
                "q [-]": st["q"],
                "v [m/s]": st["v"],
                "w [m/s]": st["w"],
                "u [m/s]": st["u"],
                "Ma [-]": st["Ma"],
                "α [deg]": st["alpha"],
                "β [deg]": st["beta"],
                "r [mm]": 1e3 * st["r"],
                "H [mm]": 1e3 * st["H"],
            }
        )

    return rows


def generate_meanline_report(
    out_left,
    out_right=None,
    left_name="Case A",
    right_name="Case B",
    filename="meanline_report.docx",
):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement, ns

    doc = Document()

    # ---------------------------------
    # Helpers
    # ---------------------------------
    def prettify_name(key):
        return key.replace("_", " ").capitalize()

    def format_value(v, scale=1.0):
        if isinstance(v, (int, float)):
            return f"{v * scale:.3f}"
        return "-" if v is None else str(v)

    def set_cell_text(cell, text, bold=False, italic=False, align="center"):
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = (
            WD_ALIGN_PARAGRAPH.LEFT if align == "left" else WD_ALIGN_PARAGRAPH.CENTER
        )
        run = p.add_run(str(text))
        run.font.name = "Times New Roman"
        run.font.size = Pt(11)
        run.bold = bold
        run.italic = italic
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = 1.0

    def set_cell_borders(cell, top=False, bottom=False):
        tc = cell._tc
        tcPr = tc.get_or_add_tcPr()

        borders = OxmlElement("w:tcBorders")

        def _border(tag):
            el = OxmlElement(tag)
            el.set(ns.qn("w:val"), "single")
            el.set(ns.qn("w:sz"), "8")
            el.set(ns.qn("w:space"), "0")
            el.set(ns.qn("w:color"), "000000")
            return el

        if top:
            borders.append(_border("w:top"))
        if bottom:
            borders.append(_border("w:bottom"))

        tcPr.append(borders)

    # ---------------------------------
    # Create table
    # ---------------------------------
    ncols = 4 if out_right else 3
    table = doc.add_table(rows=1, cols=ncols)
    table.autofit = True

    # Header
    hdr_row = table.rows[0]
    hdr = hdr_row.cells
    set_cell_text(hdr[0], "Variable", bold=True, align="left")
    set_cell_text(hdr[1], left_name, bold=True)
    if out_right:
        set_cell_text(hdr[2], right_name, bold=True)
        set_cell_text(hdr[3], "Unit", bold=True)
    else:
        set_cell_text(hdr[2], "Unit", bold=True)

    # set_row_borders(hdr_row, top=True, bottom=True)
    for cell in hdr_row.cells:
        set_cell_borders(cell, top=True, bottom=True)

    # ---------------------------------
    # Row writers
    # ---------------------------------
    def add_section(title):
        row = table.add_row()
        cells = row.cells
        set_cell_text(cells[0], title, bold=True, italic=True, align="left")
        for c in cells[1:]:
            set_cell_text(c, "")

        for cell in row.cells:
            set_cell_borders(cell, top=True, bottom=True)

        # set_row_borders(row, top=True, bottom=True)

    def add_row(name, v1, v2, unit, scale=1.0):
        row = table.add_row().cells
        set_cell_text(row[0], prettify_name(name), align="left")
        set_cell_text(row[1], format_value(v1, scale))
        if out_right:
            set_cell_text(row[2], format_value(v2, scale))
            set_cell_text(row[3], unit)
        else:
            set_cell_text(row[2], unit)

    # ---------------------------------
    # Operating conditions
    # ---------------------------------
    add_section("Operating conditions")
    oc_units = {
        "fluid_name": ("-", 1.0),
        "stage_type": ("-", 1.0),
        "inlet_property_pair": ("-", 1.0),
        "inlet_property_1": ("bar", 1e-5),
        "inlet_property_2": ("-", 1.0),
        "exit_pressure": ("bar", 1e-5),
        "mass_flow_rate": ("kg/s", 1.0),
    }

    for k, (unit, scale) in oc_units.items():
        add_row(
            k,
            out_left["inputs"].get(k),
            out_right["inputs"].get(k) if out_right else None,
            unit,
            scale,
        )

    # ---------------------------------
    # Design variables
    # ---------------------------------
    add_section("Design variables")
    dv_units = {
        "stator_inlet_angle": ("deg", 1.0),
        "stator_exit_angle": ("deg", 1.0),
        "blade_velocity_ratio": ("-", 1.0),
        "degree_reaction": ("-", 1.0),
        "radius_ratio_12": ("-", 1.0),
        "radius_ratio_23": ("-", 1.0),
        "radius_ratio_34": ("-", 1.0),
        "height_radius_ratio": ("-", 1.0),
        "zweiffel_stator": ("-", 1.0),
        "zweiffel_rotor": ("-", 1.0),
        "loss_coeff_stator": ("-", 1.0),
        "loss_coeff_rotor": ("-", 1.0),
    }

    for k, (unit, scale) in dv_units.items():
        add_row(
            k,
            out_left["inputs"].get(k),
            out_right["inputs"].get(k) if out_right else None,
            unit,
            scale,
        )

    # ---------------------------------
    # Stage performance
    # ---------------------------------
    add_section("Stage performance")
    sp_units = {
        "efficiency_tt": ("-", 1.0),
        "efficiency_ts": ("-", 1.0),
        "pressure_ratio_ts": ("-", 1.0),
        "volume_ratio_ts": ("-", 1.0),
        "flow_coefficient": ("-", 1.0),
        "work_coefficient": ("-", 1.0),
        "specific_speed": ("-", 1.0),
        "rotational_speed": ("rpm", 1.0),
        "spouting_velocity": ("m/s", 1.0),
        "power_isentropic": ("W", 1.0),
        "power_actual_tt": ("W", 1.0),
        "power_actual_ts": ("W", 1.0),
    }

    for k, (unit, scale) in sp_units.items():
        add_row(
            k,
            out_left["stage_performance"].get(k),
            out_right["stage_performance"].get(k) if out_right else None,
            unit,
            scale,
        )

    # ---------------------------------
    # Geometry
    # ---------------------------------
    for part in ("stator", "rotor"):
        add_section(f"{part.capitalize()} geometry")
        geo_units = {
            "blade_count": ("-", 1.0),
            "radius_in": ("mm", 1e3),
            "radius_out": ("mm", 1e3),
            "height": ("mm", 1e3),
            "chord": ("mm", 1e3),
            "spacing": ("mm", 1e3),
            "opening": ("mm", 1e3),
            "solidity": ("-", 1.0),
            "aspect_ratio": ("-", 1.0),
            "flaring_angle": ("deg", 1.0),
            "metal_angle_in": ("deg", 1.0),
            "metal_angle_out": ("deg", 1.0),
        }

        for k, (unit, scale) in geo_units.items():
            add_row(
                k,
                out_left["geometry"][part].get(k),
                out_right["geometry"][part].get(k) if out_right else None,
                unit,
                scale,
            )

    # Bottom rule for entire table (last row)
    last_row = table.rows[-1]
    for cell in last_row.cells:
        set_cell_borders(cell, bottom=True)

    doc.save(filename)
    return filename

